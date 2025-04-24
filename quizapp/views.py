from django.shortcuts import get_object_or_404, render, redirect
from . import models as quiz_models
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from .forms import UserRegistrationForm, QuestionForm, ChoiceFormSet
# import csv
import os
import openpyxl
from openpyxl import Workbook

# Create your views here.
@login_required(login_url="quiz:login_user")
def index(request):
    return render(request, "home.html")

@login_required(login_url="quiz:login_user")
def home(request):
    user = quiz_models.AppUser.objects.get(id=request.user.id)
    user_discipline = quiz_models.Discipline.objects.get(name=user.discipline)
    if user.has_submitted:
        return redirect("quiz:result_page")
    
    question = quiz_models.Question.objects.filter(discipline=user_discipline).order_by("id")
    total_questions= question.count()
    # choice = quiz_models.Choices.objects.filter(question=question)
    
    context = {
        "question": question,
        "total_questions": total_questions,
        "user":user
        # "choice": choice 
        
    }
    
    return render(request, "index.html", context)

def calculate_score(request):
    if request.method == "POST":
        questions = quiz_models.Question.objects.all().order_by("id")
        user = quiz_models.AppUser.objects.get(id=request.user.id)

        score = 0
        for question in questions:
            selected_choice_id = request.POST.get(f"choice_{question.id}")
            if selected_choice_id:
                try:
                    selected_choice = quiz_models.Choices.objects.get(id=selected_choice_id)
                    
                    # save user's answers
                    quiz_models.UserAnswer.objects.update_or_create(
                        user=user,
                        question=question,
                        defaults={"selected_choice": selected_choice}
                    )
                    if selected_choice.is_correct:
                        score += 2
                except quiz_models.Choices.DoesNotExist:
                    pass

        # Save score to user
        user.score = score
        user.has_submitted = True
        user.save()

        # Write to Excel file
        xlsx_path = os.path.join("plugin", "results.xlsx")
        if os.path.exists(xlsx_path):
            wb = openpyxl.load_workbook(xlsx_path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(["Full Name", "Score", "Discipline"])  # Header

        # Check if the user is already in the sheet
        user_already_written = any(
            row[0].value == user.get_full_name() for row in ws.iter_rows(min_row=3)
        )

        if not user_already_written:
            ws.append([user.get_full_name(), user.score, str(user.discipline)])
            wb.save(xlsx_path)

    return redirect("quiz:result_page")

@login_required(login_url="quiz:login_user")
def result_page(request):
    user_id = quiz_models.AppUser.objects.get(id=request.user.id)
    discipline = quiz_models.Discipline.objects.get(name=user_id.discipline)
    questions = quiz_models.Question.objects.filter(discipline=discipline)
    user_score = user_id.score
    # correct_answers = user_id.num_of_correct_choices
    len_of_questions = questions.count()
    correct_count = quiz_models.Choices.objects.filter(is_correct=True).count()
    
    
    context = {
        "user_score":user_score,
        # "correct_answers":correct_answers,
        "len_of_questions":len_of_questions,
        # "correct_count":correct_count,
        
    }
    return render(request, 'result.html', context)
                    
def login_user(request):
    if request.user.is_authenticated:
        return redirect("quiz:index")
    
    if request.method == "POST":
        username = request.POST.get("username")
        password = request.POST.get("pswd")
        
        my_user = authenticate(username=username, password=password)
        
        if my_user is not None:
            login(request, my_user)
            messages.success(request, "Login Successful!!!")
            return redirect(to="quiz:index")
        else:
            messages.error(request, "There was a problem logging you in\nCheck login credentials")
            return redirect(to="quiz:login_user")
        
    return render(request, 'login.html')

def register_user(request):
    form = UserRegistrationForm()
    if request.method == "POST":
        form = UserRegistrationForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Sign Up success")
            return redirect("quiz:login_user")
        messages.error(request, "There was a problem!!")
        return redirect("quiz:register_user")
    
    context = {
        "form":form
    }
    return render(request, "register.html", context)

@login_required
def add_question(request):
    try:
        # Check if user is an instructor
        user = quiz_models.AppUser.objects.get(is_instructor=True, id=request.user.id)
        
        if request.method == 'POST':
            form = QuestionForm(request.POST)
            formset = ChoiceFormSet(request.POST)

            if form.is_valid() and formset.is_valid():
                # Save the question but don't commit yet
                question = form.save(commit=False)
                # Set the instructor to the current user
                question.instructor = request.user
                question.save()
                
                # Now save the choices
                choices = formset.save(commit=False)
                for choice in choices:
                    choice.question = question
                    choice.save()
                
                messages.success(request, "Question added successfully!")
                return redirect('quiz:index')

        else:
            # Initialize empty forms
            form = QuestionForm(initial={'instructor': request.user})
            formset = ChoiceFormSet()

    except quiz_models.AppUser.DoesNotExist:
        messages.warning(request, "You are not allowed here!!!")
        return redirect(to="quiz:index")

    return render(request, 'add_question.html', {'form': form, 'formset': formset})


def question_list(request, pk):
    user = quiz_models.AppUser.objects.get(id=pk)
    questions = quiz_models.Question.objects.filter(instructor=user)
    
    context = {
        "questions": questions
    }
    return render(request, "questions.html", context)

@login_required
def edit_question(request, question_id):
    try:
        # Check if user is an instructor
        user = quiz_models.AppUser.objects.get(is_instructor=True, id=request.user.id)
        
        # Get the question to edit
        question = get_object_or_404(quiz_models.Question, id=question_id, instructor=request.user)
        
        if request.method == 'POST':
            form = QuestionForm(request.POST, instance=question)
            formset = ChoiceFormSet(request.POST, instance=question)

            if form.is_valid() and formset.is_valid():
                form.save()
                instances = formset.save(commit=False)
                
                # Delete any marked for deletion
                for obj in formset.deleted_objects:
                    obj.delete()
                
                # Save new and updated choices
                for instance in instances:
                    instance.question = question
                    instance.save()
                
                messages.success(request, "Question updated successfully!")
                return redirect('quiz:index')

        else:
            form = QuestionForm(instance=question)
            # Initialize formset with exactly 4 choices (existing + empty)
            formset = ChoiceFormSet(instance=question, queryset=question.choices_set.all().order_by('id'))
            
            # If there are fewer than 4 choices, adjust the extra forms
            existing_choices = question.choices_set.count()
            if existing_choices < 4:
                formset.extra = 4 - existing_choices
            else:
                formset.extra = 0

    except quiz_models.AppUser.DoesNotExist:
        messages.warning(request, "You are not allowed to edit questions!")
        return redirect(to="quiz:index")

    return render(request, 'edit_question.html', {
        'form': form,
        'formset': formset,
        'question': question
    })